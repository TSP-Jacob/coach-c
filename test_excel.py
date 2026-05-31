import openpyxl

wb = openpyxl.load_workbook('agents.xlsx')
gatineau = wb['Gatineau']
print(f"Tables in Gatineau: {gatineau.tables.keys()}")
if gatineau.tables:
    for tbl in gatineau.tables.values():
        print(f"Table name: {tbl.displayName}, Ref: {tbl.ref}")

# Let's also see what columns actually have data
print("Row 2 data:")
print([cell.value for cell in gatineau[2]])
